# -*- coding: utf-8 -*-
"""Rebuild Excel workbook: configurable environments + storage tiers (hot/warm/cold/frozen) + cloud/on-prem."""
from pathlib import Path
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(r"C:\Users\DK\Mygithub\elastic-cloud-cost-allocator")
OUT = ROOT / "elastic-cloud-cost-allocator.xlsx"
SEED = json.loads((ROOT / "docs" / "seed-measurement.json").read_text(encoding="utf-8"))

YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEADER = PatternFill("solid", fgColor="E8F0EF")
TITLE = Font(bold=True, size=14, color="1F6B66")
BOLD = Font(bold=True)
THIN = Border(
    left=Side(style="thin", color="D9E2DF"),
    right=Side(style="thin", color="D9E2DF"),
    top=Side(style="thin", color="D9E2DF"),
    bottom=Side(style="thin", color="D9E2DF"),
)

# Sample topology: hot+frozen enabled (warm/cold off). Toggle Tiers!Enabled and zero unused rates.
TIERS = [
    # code, label, enabled
    ("hot", "Hot", True),
    ("warm", "Warm", False),
    ("cold", "Cold", False),
    ("frozen", "Frozen", True),
]

# code, label, applySurcharge, hot, warm, cold, frozen, kibana, integrations, tiebreaker, snapshots
ENVS = [
    ("PRE", "Pre-production (sample)", True, 3.5, 0, 0, 0.35, 0.28, 0.14, 0, 200),
    ("PRO", "Production (sample)", True, 5.0, 0, 0, 1.0, 0.56, 0.14, 0.07, 500),
    ("MON", "Monitoring dedicated (sample)", False, 0.5, 0, 0, 0, 0, 0, 0, 10),
]

wb = Workbook()

# --- README ---
wr = wb.active
wr.title = "README"
wr["A1"] = "Elastic Cost Allocator — Cloud & on-prem · configurable tiers & environments"
wr["A1"].font = TITLE
wr["A3"] = "1. Tiers: enable only the storage tiers you use (hot-only, +warm/cold, +frozen). Sample ships with hot+frozen."
wr["A4"] = "2. Global: set Deployment (cloud|onprem), cost unit (ECU/EUR/USD), hours, annual commit/budget, margin, snapshot alloc tier."
wr["A5"] = "3. Environments: add/rename codes (DEV, INT, UAT, PROD, MONITORING, …). Set Apply surcharge and per-tier rates."
wr["A6"] = "4. Measurement: one row per index family. Fill GB columns for enabled tiers. Env must match Environments!Code."
wr["A7"] = "5. Read Without / With Admin Cost and Any Family. HTML twin: web/index.html"
wr["A9"] = "On-prem: set Deployment=onprem, unit=EUR (or your currency), zero Cloud-only platform rates if unused."
wr["A10"] = "https://github.com/superlopezg/elastic-cloud-cost-allocator"
wr.column_dimensions["A"].width = 120

# --- Tiers ---
wt = wb.create_sheet("Tiers")
wt["A1"] = "Storage tiers — set Enabled=TRUE only for tiers present in your cluster(s)."
wt["A1"].font = TITLE
wt.merge_cells("A1:C1")
for c, h in enumerate(["Code", "Label", "Enabled"], 1):
    cell = wt.cell(3, c, h)
    cell.fill = HEADER
    cell.font = BOLD
for i, row in enumerate(TIERS, 4):
    for c, v in enumerate(row, 1):
        cell = wt.cell(i, c, v)
        cell.fill = YELLOW
        cell.border = THIN
wt.column_dimensions["A"].width = 12
wt.column_dimensions["B"].width = 16
wt.column_dimensions["C"].width = 12
wt["A9"] = "Tip: hot-only → enable Hot only. ILM with warm/cold → enable those and enter GB + rates. Frozen searchables → enable Frozen."
wt.merge_cells("A9:C9")

# --- Environments ---
ws = wb.create_sheet("Environments")
ws["A1"] = "Environments — configure as many clusters as you need. Rates for disabled tiers can stay 0."
ws["A1"].font = TITLE
ws.merge_cells("A1:K1")
ws["A2"] = "Yellow = editable. Measurement.env must match Code. Surcharge=TRUE → commit/budget uplift; FALSE → dedicated."
headers = [
    "Code", "Label", "Apply surcharge",
    "Hot /h", "Warm /h", "Cold /h", "Frozen /h",
    "Kibana /h", "Integrations /h", "Tiebreaker /h", "Snapshots /month",
]
for c, h in enumerate(headers, 1):
    cell = ws.cell(4, c, h)
    cell.fill = HEADER
    cell.font = BOLD
for i, row in enumerate(ENVS, 5):
    for c, v in enumerate(row, 1):
        cell = ws.cell(i, c, v)
        cell.fill = YELLOW
        cell.border = THIN
ws.column_dimensions["A"].width = 12
ws.column_dimensions["B"].width = 32
for col in range(3, 12):
    ws.column_dimensions[get_column_letter(col)].width = 14

# --- Global ---
wg = wb.create_sheet("Global")
wg["A1"] = "Global contract / budget parameters"
wg["A1"].font = TITLE
wg["A3"] = "Deployment (cloud | onprem)"
wg["B3"] = "cloud"
wg["B3"].fill = YELLOW
wg["A4"] = "Cost unit (ECU, EUR, USD, …)"
wg["B4"] = "ECU"
wg["B4"].fill = YELLOW
wg["A5"] = "Hours / month"
wg["B5"] = 730
wg["B5"].fill = YELLOW
wg["A6"] = "Annual commit / budget"
wg["B6"] = 100000
wg["B6"].fill = YELLOW
wg["A7"] = "Margin on commit"
wg["B7"] = 0.03
wg["B7"].fill = YELLOW
wg["A8"] = "Other cost / year"
wg["B8"] = 500
wg["B8"].fill = YELLOW
wg["A9"] = "Snapshot alloc tier (auto | hot | warm | cold | frozen)"
wg["B9"] = "auto"
wg["B9"].fill = YELLOW
wg["A11"] = "Hours / year"
wg["B11"] = "=B5*12"
wg["A12"] = "Annual target"
wg["B12"] = "=B6*(1+B7)"
# Tier rate sum columns D:G; platform H:J; snapshots K
wg["A14"] = "Surchargeable base / year (Apply surcharge = TRUE)"
wg["B14"] = (
    '=SUMPRODUCT((Environments!C5:C50=TRUE)*(Environments!D5:D50+Environments!E5:E50'
    '+Environments!F5:F50+Environments!G5:G50)*$B$11)'
    '+SUMPRODUCT((Environments!C5:C50=TRUE)*Environments!K5:K50*12)'
)
wg["A15"] = "Dedicated envs / year (Apply surcharge = FALSE)"
wg["B15"] = (
    '=SUMPRODUCT((Environments!C5:C50=FALSE)*(Environments!D5:D50+Environments!E5:E50'
    '+Environments!F5:F50+Environments!G5:G50)*$B$11)'
    '+SUMPRODUCT((Environments!C5:C50=FALSE)*Environments!K5:K50*12)'
)
wg["A16"] = "Platform annual (Kibana+Integrations+Tiebreaker on surchargeable envs)"
wg["B16"] = (
    '=SUMPRODUCT((Environments!C5:C50=TRUE)*(Environments!H5:H50+Environments!I5:I50'
    '+Environments!J5:J50)*$B$11)'
)
wg["A17"] = "Computed surcharge"
wg["B17"] = "=IF(B14=0,0,(B12-B15)/B14-1)"
wg["A18"] = "Estimated consumption"
wg["B18"] = "=B14+B15+B16+B8"
wg.column_dimensions["A"].width = 72
wg.column_dimensions["B"].width = 18

# --- Measurement ---
# A Family, B Env, C What, D eps15, E kbdoc, F docs, G hot, H warm, I cold, J frozen, K indices, L epsLife, M age
wm = wb.create_sheet("Measurement")
wm["A1"] = "Measurement — GB columns for all tiers. Leave warm/cold at 0 if those tiers are disabled."
wm["A1"].font = TITLE
wm.merge_cells("A1:M1")
mheaders = [
    "Family", "Env", "What", "Events/s 15m", "KB/doc", "Docs",
    "GB hot", "GB warm", "GB cold", "GB frozen",
    "Indices", "Events/s life", "Age days",
]
for c, h in enumerate(mheaders, 1):
    cell = wm.cell(2, c, h)
    cell.fill = HEADER
    cell.font = BOLD
for i, r in enumerate(SEED, 3):
    vals = [
        r["family"], r["env"], r.get("what", ""), r["eps15"], r["kbdoc"], r["docs"],
        r["gbHot"], 0, 0, r["gbFrozen"],
        r["indices"], r["epsLife"], r["ageDays"],
    ]
    for c, v in enumerate(vals, 1):
        cell = wm.cell(i, c, v)
        if c in (1, 2, 3, 7, 8, 9, 10):
            cell.fill = YELLOW
last_meas = 2 + len(SEED)
dv = DataValidation(type="list", formula1="=Environments!$A$5:$A$50", allow_blank=True)
wm.add_data_validation(dv)
dv.add(f"B3:B{last_meas + 50}")
for col, w in enumerate([28, 10, 40, 12, 10, 14, 10, 10, 10, 12, 10, 12, 10], 1):
    wm.column_dimensions[get_column_letter(col)].width = w

# Allocation: shares for all 4 tiers; capacity sums all; snapshots use frozen share with hot fallback (auto-like)
def build_alloc(sheet_name, with_admin):
    wa = wb.create_sheet(sheet_name)
    wa["A1"] = sheet_name + (
        " — surcharge when Environments!Apply surcharge is TRUE" if with_admin else " — no surcharge"
    )
    wa["A1"].font = TITLE
    wa["A2"] = "Surcharge rate"
    wa["B2"] = "=Global!B17" if with_admin else 0
    if with_admin:
        wa["B2"].fill = YELLOW
    headers = [
        "Family", "Env",
        "GB hot", "Share hot", "GB warm", "Share warm", "GB cold", "Share cold", "GB frozen", "Share frozen",
        "Hot /h", "Warm /h", "Cold /h", "Frozen /h", "Snapshots/mo",
        "Cost/mo capacity", "Cost/mo snapshots", "Cost/mo Elastic", "Cost/mo surcharge", "Cost/mo total", "Cost/year",
    ]
    for c, h in enumerate(headers, 1):
        cell = wa.cell(4, c, h)
        cell.fill = HEADER
        cell.font = BOLD
    n = len(SEED)
    for i in range(n):
        r = 5 + i
        m = 3 + i
        wa.cell(r, 1, f"=Measurement!A{m}")
        wa.cell(r, 2, f"=Measurement!B{m}")
        # GB + share for hot(G), warm(H), cold(I), frozen(J)
        # cols: 3=GBhot 4=shareHot 5=GBwarm 6=shareWarm 7=GBcold 8=shareCold 9=GBfrz 10=shareFrz
        pairs = [(3, 4, "G"), (5, 6, "H"), (7, 8, "I"), (9, 10, "J")]
        for gb_c, sh_c, mcol in pairs:
            wa.cell(r, gb_c, f"=Measurement!{mcol}{m}")
            wa.cell(
                r,
                sh_c,
                f'=IF(SUMIF(Measurement!$B:$B,B{r},Measurement!${mcol}:${mcol})=0,0,'
                f'{get_column_letter(gb_c)}{r}/SUMIF(Measurement!$B:$B,B{r},Measurement!${mcol}:${mcol}))',
            )
        # rates: Environments D=4 hot, E=5 warm, F=6 cold, G=7 frozen, K=11 snapshots
        wa.cell(r, 11, f'=IFERROR(VLOOKUP(B{r},Environments!$A$5:$K$50,4,FALSE),0)')
        wa.cell(r, 12, f'=IFERROR(VLOOKUP(B{r},Environments!$A$5:$K$50,5,FALSE),0)')
        wa.cell(r, 13, f'=IFERROR(VLOOKUP(B{r},Environments!$A$5:$K$50,6,FALSE),0)')
        wa.cell(r, 14, f'=IFERROR(VLOOKUP(B{r},Environments!$A$5:$K$50,7,FALSE),0)')
        wa.cell(r, 15, f'=IFERROR(VLOOKUP(B{r},Environments!$A$5:$K$50,11,FALSE),0)')
        # capacity = sum(share*rate)*hours
        wa.cell(r, 16, f"=(D{r}*K{r}+F{r}*L{r}+H{r}*M{r}+J{r}*N{r})*Global!$B$5")
        # snapshots: frozen share if env has frozen GB else hot (auto)
        wa.cell(
            r,
            17,
            f"=IF(SUMIF(Measurement!$B:$B,B{r},Measurement!$J:$J)=0,D{r}*O{r},J{r}*O{r})",
        )
        wa.cell(r, 18, f"=P{r}+Q{r}")
        if with_admin:
            wa.cell(r, 19, f'=IF(IFERROR(VLOOKUP(B{r},Environments!$A$5:$K$50,3,FALSE),FALSE),R{r}*$B$2,0)')
        else:
            wa.cell(r, 19, 0)
        wa.cell(r, 20, f"=R{r}+S{r}")
        wa.cell(r, 21, f"=T{r}*12")
    for col in range(1, 22):
        wa.column_dimensions[get_column_letter(col)].width = 12
    wa.column_dimensions["A"].width = 28

build_alloc("Without Admin Cost", False)
build_alloc("With Admin Cost", True)

# --- Any Family ---
wy = wb.create_sheet("Any Family")
wy["A1"] = "Pick any Measurement family — rates from Environments via VLOOKUP (any env code)."
wy["A1"].font = TITLE
wy["A3"] = "Family"
wy["B3"] = SEED[0]["family"] if SEED else ""
wy["B3"].fill = YELLOW
wy["A4"] = "Env"
wy["B4"] = '=IFERROR(INDEX(Measurement!B:B,MATCH(B3,Measurement!A:A,0)),"")'
row = 5
for label, mcol in [("GB hot", "G"), ("GB warm", "H"), ("GB cold", "I"), ("GB frozen", "J")]:
    wy.cell(row, 1, label)
    wy.cell(row, 2, f'=IFERROR(SUMIF(Measurement!A:A,B3,Measurement!{mcol}:{mcol}),0)')
    row += 1
for label, mcol, gb_row in [
    ("Hot share", "G", 5),
    ("Warm share", "H", 6),
    ("Cold share", "I", 7),
    ("Frozen share", "J", 8),
]:
    wy.cell(row, 1, label)
    wy.cell(
        row,
        2,
        f'=IF(SUMIF(Measurement!B:B,B4,Measurement!{mcol}:{mcol})=0,0,B{gb_row}/SUMIF(Measurement!B:B,B4,Measurement!{mcol}:{mcol}))',
    )
    row += 1
# rates
for label, col_idx in [("Hot /h", 4), ("Warm /h", 5), ("Cold /h", 6), ("Frozen /h", 7), ("Snapshots / mo", 11)]:
    wy.cell(row, 1, label)
    wy.cell(row, 2, f'=IFERROR(VLOOKUP(B4,Environments!$A$5:$K$50,{col_idx},FALSE),0)')
    row += 1
wy.cell(row, 1, "Apply surcharge?")
wy.cell(row, 2, '=IFERROR(VLOOKUP(B4,Environments!$A$5:$K$50,3,FALSE),FALSE)')
row += 1
# B5..B8 GB, B9..B12 shares, B13..B16 rates, B17 snaps, B18 surcharge flag
wy.cell(row, 1, "Cost/mo capacity")
wy.cell(row, 2, "=(B9*B13+B10*B14+B11*B15+B12*B16)*Global!B5")
row += 1
wy.cell(row, 1, "Cost/mo snapshots")
wy.cell(row, 2, "=IF(SUMIF(Measurement!B:B,B4,Measurement!J:J)=0,B9*B17,B12*B17)")
row += 1
wy.cell(row, 1, "Cost/mo Elastic")
wy.cell(row, 2, f"=B{row-2}+B{row-1}")
cap_row, snap_row, elas_row = row - 2, row - 1, row
row += 1
wy.cell(row, 1, "Cost/mo surcharge")
wy.cell(row, 2, f"=IF(B18,B{elas_row}*Global!B17,0)")
sur_row = row
row += 1
wy.cell(row, 1, "Cost/mo total")
wy.cell(row, 2, f"=B{elas_row}+B{sur_row}")
tot_row = row
row += 1
wy.cell(row, 1, "Cost/year")
wy.cell(row, 2, f"=B{tot_row}*12")
wy.column_dimensions["A"].width = 22
wy.column_dimensions["B"].width = 36

wb.save(OUT)
print("Wrote", OUT, "sheets", wb.sheetnames, "meas rows", len(SEED))
